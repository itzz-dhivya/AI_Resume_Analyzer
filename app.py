import streamlit as st
import os
import hashlib
import pandas as pd
import mysql.connector
from mysql.connector import Error
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
import tempfile
import requests
import time
from streamlit_lottie import st_lottie
import plotly.express as px

# ----------------- CONFIG -----------------
st.set_page_config(page_title="AI Resume Analyzer", layout="centered", initial_sidebar_state="auto")

# ----------------- DATABASE -----------------
DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "Dhivya",
    "database": "resume_analyzer"
}

def create_conn():
    try:
        conn = mysql.connector.connect(**DB_CONFIG)
        return conn
    except Error as e:
        st.error(f"Database connection failed: {e}")
        return None

# ----------------- HASHING -----------------
SALT = "a_random_salt_123!@#"

def hash_password(password: str) -> str:
    return hashlib.sha256((SALT + password).encode()).hexdigest()

def check_password(password: str, hashed: str) -> bool:
    return hash_password(password) == hashed

# ----------------- STYLES & BG -----------------
def set_bg_and_styles():
    st.markdown("""
    <style>
    .stApp {
        background: linear-gradient(135deg, #0f172a 0%, #0f172a 40%, #0b3a66 100%);
        color: #e6eef8;
    }
    .glass-card {
        background: rgba(255, 255, 255, 0.06);
        backdrop-filter: blur(6px);
        border-radius: 12px;
        padding: 18px;
        box-shadow: 0 8px 30px rgba(2,6,23,0.6);
        border: 1px solid rgba(255,255,255,0.06);
        margin-bottom: 16px;
    }
    .btn-primary > button {
        background: linear-gradient(90deg, rgba(255,255,255,0.06), rgba(255,255,255,0.04));
        color: white;
        border-radius: 10px;
        padding: 8px 14px;
    }
    .small-note { font-size:12px; color:#c9d7f0; }
    </style>
    """, unsafe_allow_html=True)

set_bg_and_styles()

def card_container(content):
    st.markdown(f"<div class='glass-card'>{content}</div>", unsafe_allow_html=True)

def load_lottie(url):
    try:
        r = requests.get(url, timeout=6)
        if r.status_code == 200:
            return r.json()
    except Exception:
        return None
    return None

# ----------------- ROLE SKILLS & COURSES -----------------
role_skills_courses = {
    "Software Developer": {"skills": ["Python","Java","C++"], "courses":["https://www.freecodecamp.org/learn","https://www.coursera.org/professional-certificates/google-it-automation"]},
    "Web Developer": {"skills":["HTML","CSS","JavaScript"], "courses":["https://www.freecodecamp.org/learn/responsive-web-design/","https://www.codecademy.com/learn/paths/web-development"]},
    "Mobile App Developer": {"skills":["Flutter","React Native","Swift","Kotlin"], "courses":["https://flutter.dev/docs/codelabs","https://www.codecademy.com/learn/paths/mobile-development"]},
    "Game Developer": {"skills":["Unity","C#","Unreal Engine"], "courses":["https://learn.unity.com/","https://www.coursera.org/learn/game-design-and-development"]},
    "Embedded Systems Developer": {"skills":["C","Microcontrollers","RTOS"], "courses":["https://www.edx.org/course/embedded-systems","https://www.coursera.org/specializations/embedded-software"]},
    "Python Developer": {"skills":["Python"], "courses":["https://www.freecodecamp.org/learn/python/"]},
    "Java Developer": {"skills":["Java"], "courses":["https://www.coursera.org/specializations/java-programming"]},
    "C++ Developer": {"skills":["C++"], "courses":["https://www.coursera.org/learn/c-plus-plus-a"]},
    "Data Analyst": {"skills":["Excel","SQL","Python","Tableau"], "courses":["https://www.coursera.org/professional-certificates/google-data-analytics","https://www.freecodecamp.org/learn/data-analysis-with-python/"]},
    "Data Scientist": {"skills":["Python","SQL","Pandas","NumPy"], "courses":["https://www.coursera.org/professional-certificates/ibm-data-science"]}
}

# ----------------- SKILL EXTRACTION -----------------
def extract_skills(text):
    keywords = []
    for role_info in role_skills_courses.values():
        for skill in role_info.get("skills", []):
            if skill.lower() in text.lower() and skill not in keywords:
                keywords.append(skill)
    return keywords

def extract_skills_from_pdf(pdf_path):
    try:
        reader = PdfReader(pdf_path)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + " "
        return extract_skills(text)
    except Exception:
        return []

def extract_skills_from_docx(docx_path):
    try:
        doc = Document(docx_path)
        text = " ".join([para.text for para in doc.paragraphs])
        return extract_skills(text)
    except Exception:
        return []

# ----------------- PDF REPORT -----------------
def generate_pdf(profile, missing_skills, missing_courses):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "AI Resume Analyzer Report", ln=True, align="C")
    pdf.ln(8)
    pdf.set_font("Arial", "", 12)
    pdf.cell(0, 8, f"Name: {profile.get('name')}", ln=True)
    pdf.cell(0, 8, f"Career Goal: {profile.get('career_goal')}", ln=True)
    pdf.cell(0, 8, f"Skills: {profile.get('skills')}", ln=True)
    pdf.ln(6)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Missing Skills & Suggested Courses:", ln=True)
    pdf.set_font("Arial", "", 11)
    for i, skill in enumerate(missing_skills):
        course = missing_courses[i] if i < len(missing_courses) else "No course available"
        pdf.multi_cell(0, 8, f"- {skill}: {course}")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp_path = tmp.name
    tmp.close()
    pdf.output(tmp_path)
    with open(tmp_path, "rb") as f:
        data = f.read()
    os.remove(tmp_path)
    return data

# ----------------- NAVIGATION -----------------
def go_to(page_name):
    st.session_state["page"] = page_name
    

if "page" not in st.session_state:
    st.session_state["page"] = "home"
if "user" not in st.session_state:
    st.session_state["user"] = None

# ----------------- MODULES -----------------
def home_module():
    lottie_home = load_lottie("https://assets10.lottiefiles.com/packages/lf20_jcikwtux.json")
    col1, col2 = st.columns([2,1])
    with col1:
        if lottie_home:
            st_lottie(lottie_home, height=220)
        card_container("""
            <h1>👋 Welcome to <strong>AI Resume Analyzer</strong></h1>
            <p class='small-note'>Upload your resume (PDF/DOCX) to get role suggestions, missing skills and curated courses.</p>
        """)
        st.markdown("<div class='btn-primary'>", unsafe_allow_html=True)
        c1, c2, c3 = st.columns([1,1,1])
        with c1:
            if st.button("Login"):
                go_to("login")
        with c2:
            if st.button("Register"):
                go_to("register")
    with col2:
        st.markdown("<div class='glass-card' style='text-align:left'>", unsafe_allow_html=True)
        st.subheader("Quick Tips for Best Results ✅")
        st.write("- Supported file types: **PDF** and **DOCX** only.")
        st.write("- Make sure your resume highlights **projects, internships, and certifications**.")
        st.write("- Add a short **career goal statement** for personalized recommendation")
        st.write("- Main:Only for computer Sector students")
        st.markdown("</div>", unsafe_allow_html=True)

def login_module():
    card_container("<h2>🔐 Login</h2>")
    username = st.text_input("Username", key="login_username")
    password = st.text_input("Password", type="password", key="login_password")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("Login", key="login_btn"):
            if not username or not password:
                st.warning("Enter username & password.")
            else:
                conn = create_conn()
                if conn:
                    try:
                        cursor = conn.cursor(dictionary=True)
                        cursor.execute("SELECT * FROM users WHERE username=%s", (username,))
                        user = cursor.fetchone()
                        conn.close()
                        if user and check_password(password, user.get("password_hash")):
                            st.session_state["user"] = {"id": user["id"], "username": user["username"]}
                            st.success("Login successful ✅")
                            go_to("resume_upload")
                        else:
                            st.error("Invalid credentials ✖")
                    except Exception as e:
                        st.error(f"Login error: {e}")
                else:
                    st.error("DB connection failed.")
    with col2:
        if st.button("Back to Home"):
            go_to("home")

def register_module():
    card_container("<h2>📝 Create Account</h2>")
    new_user = st.text_input("Username", key="reg_username")
    new_email = st.text_input("Email", key="reg_email")
    new_pass = st.text_input("Password", type="password", key="reg_password")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("Register", key="reg_btn"):
            if not new_user.strip() or not new_pass.strip() or not new_email.strip():
                st.warning("Please fill all fields including email.")
            else:
                conn = create_conn()
                if conn:
                    try:
                        cursor = conn.cursor()
                        cursor.execute(
                            "SELECT id FROM users WHERE username=%s OR email=%s",
                            (new_user, new_email)
                        )
                        if cursor.fetchone():
                            st.error("Username or email already exists.")
                        else:
                            hashed = hash_password(new_pass)
                            cursor.execute(
                                "INSERT INTO users(username, email, password_hash) VALUES(%s, %s, %s)",
                                (new_user, new_email, hashed)
                            )
                            conn.commit()
                            st.success("Registration successful 🎉 — please login.")
                            st.session_state["reg_username"] = ""
                            st.session_state["reg_email"] = ""
                            st.session_state["reg_password"] = ""
                            go_to("login")
                    except Exception as e:
                        st.error(f"Registration error: {e}")
                    finally:
                        conn.close()
    with col2:
        if st.button("Back to Home"):
            go_to("home")

# ----------------- RESUME UPLOAD MODULE -----------------
def resume_upload_module():
    card_container("<h2>📄 Upload Resume</h2>")
    uploaded_file = st.file_uploader("Upload Resume (PDF/DOCX)", type=["pdf","docx"])
    if uploaded_file:
        save_dir = "data/sample_resumes"
        os.makedirs(save_dir, exist_ok=True)
        save_path = os.path.join(save_dir, f"{int(time.time())}_{uploaded_file.name}")
        with open(save_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
        st.success("Resume uploaded ✅")
        st.session_state["resume_path"] = save_path

        if uploaded_file.name.lower().endswith(".pdf"):
            resume_skills = extract_skills_from_pdf(save_path)
        else:
            resume_skills = extract_skills_from_docx(save_path)

        if not resume_skills:
            st.warning("No recognized skills found. Try using a different resume or add skills manually.")
        else:
            st.session_state["user_skills"] = resume_skills
            st.success(f"Detected Skills: {', '.join(resume_skills)}")
        st.markdown("---")
        col1, col2 = st.columns(2)
        with col1:
            if st.button("Next: Recommendations") and resume_skills:
                go_to("recommendation")
        with col2:
            if st.button("Back"):
                go_to("home")
    else:
        st.info("Upload a PDF or DOCX resume to start.")
        if st.button("Back"):
            go_to("home")

# ----------------- RECOMMENDATION MODULE -----------------
def recommendation_module():
    resume_skills = st.session_state.get("user_skills", [])
    card_container("<h2>🔎 Recommendations</h2>")
    recommendations = []

    for role, info in role_skills_courses.items():
        role_skills = info.get("skills", [])
        matched_skills = [s for s in role_skills if s in resume_skills]
        if matched_skills:
            missing_skills = [s for s in role_skills if s not in resume_skills]
            missing_courses = info.get("courses", ["No course available"] * len(missing_skills))
            recommendations.append({
                "role": role,
                "matched_skills": matched_skills,
                "missing_skills": missing_skills,
                "missing_courses": missing_courses
            })

    if not recommendations:
        st.warning("No suitable role found. Consider adding skills to your resume.")
    else:
        for rec in recommendations:
            st.markdown(f"#### 🎯 {rec['role']}")
            st.markdown(f"- ✅ Matched Skills: {', '.join(rec['matched_skills'])}")
            if rec["missing_skills"]:
                for i, skill in enumerate(rec["missing_skills"]):
                    course = rec["missing_courses"][i] if i < len(rec["missing_courses"]) else "No course available"
                    st.markdown(f"- ⚡ Missing Skill: **{skill}** → [Course]({course})")
            st.markdown("---")

    col1, col2, col3 = st.columns([1,1,1])
    with col1:
        if st.button("Download Skill Report PDF") and recommendations:
            profile = {
                "name": st.session_state.get("user", {}).get("username", "N/A"),
                "career_goal": "Based on Resume",
                "skills": ", ".join(resume_skills)
            }
            pdf_bytes = generate_pdf(profile, recommendations[0]["missing_skills"], recommendations[0]["missing_courses"])
            st.download_button("Download PDF", pdf_bytes, file_name="skill_report.pdf", mime="application/pdf")
    with col2:
        if st.button("Next: Skill Chart") and recommendations:
            go_to("chart_page")
    with col3:
        if st.button("Back"):
            go_to("resume_upload")

# ----------------- CHART MODULE -----------------
def chart_module():
    resume_skills = st.session_state.get("user_skills", [])
    card_container("<h2>📊 Skill Match Chart</h2>")
    roles_chart_data = []
    for role, info in role_skills_courses.items():
        role_skills = info.get("skills", [])
        matched = len([s for s in role_skills if s in resume_skills])
        missing = max(0, len(role_skills) - matched)
        if matched > 0:
            roles_chart_data.append({"role": role, "Matched": matched, "Missing": missing})

    if roles_chart_data:
        df_chart = pd.DataFrame(roles_chart_data)
        df_chart = df_chart.melt(id_vars="role", value_vars=["Matched","Missing"], var_name="Status", value_name="Count")
        fig = px.bar(df_chart, x="role", y="Count", color="Status", barmode="group", title="Skills Match Chart per Role", height=450)
        st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("Upload resume or use demo to see chart")

    c1, c2 = st.columns(2)
    with c1:
        if st.button("Finish"):
            go_to("thankyou")
    with c2:
        if st.button("Back"):
            go_to("recommendation")

# ----------------- THANK YOU MODULE -----------------
def thankyou_module():
    name = st.session_state.get("user", {}).get("username", "User")
    lottie_party = load_lottie("https://assets5.lottiefiles.com/packages/lf20_jbrw3hcz.json")
    if lottie_party:
        st_lottie(lottie_party, height=220)
    st.success(f"🎉 Thank you, {name}, for using AI Resume Analyzer!")
    st.markdown("You can logout below.")
    if st.button("Logout"):
        for k in list(st.session_state.keys()):
            if k != "page":
                st.session_state.pop(k)
        st.session_state["page"] = "home"
    

# ----------------- PAGE ROUTING -----------------
page = st.session_state.get("page", "home")
if page == "home":
    home_module()
elif page == "login":
    login_module()
elif page == "register":
    register_module()
elif page == "resume_upload":
    resume_upload_module()
elif page == "recommendation":
    recommendation_module()
elif page == "chart_page":
    chart_module()
elif page == "thankyou":
    thankyou_module()
else:
    st.error("Unknown page. Returning home.")
    st.session_state["page"] = "home"
    home_module()
